import re
import zipfile
from docx import Document

path = r"C:\Users\HP\Desktop\Side Hustle\Aksen-Labs\Aksen-Labs-Business-Blueprint-v0.2.docx"
doc = Document(path)
with zipfile.ZipFile(path) as archive:
    xml = archive.read("word/document.xml").decode("utf-8")

word_count = sum(len(p.text.split()) for p in doc.paragraphs)
word_count += sum(len(cell.text.split()) for table in doc.tables for row in table.rows for cell in row.cells)

checks = {
    "paragraphs": len(doc.paragraphs),
    "tables": len(doc.tables),
    "inline_images": len(doc.inline_shapes),
    "words": word_count,
    "heading_1": sum(1 for p in doc.paragraphs if p.style.name == "Heading 1"),
    "heading_2": sum(1 for p in doc.paragraphs if p.style.name == "Heading 2"),
    "heading_3": sum(1 for p in doc.paragraphs if p.style.name == "Heading 3"),
    "hyperlinks": xml.count("<w:hyperlink"),
    "numbered_paragraphs": xml.count("<w:numPr>"),
    "page_breaks": xml.count('w:type="page"'),
    "alt_text_records": xml.count('descr="'),
    "header_rows": xml.count("tblHeader"),
    "legacy_gold_absent": "F4C542" not in xml,
    "legacy_cyan_absent": "38D6D6" not in xml,
    "signal_green_present": "7CFF62" in xml,
    "document_xml_ok": bool(re.search(r"<w:body>", xml)),
}
for key, value in checks.items():
    print(f"{key}: {value}")
